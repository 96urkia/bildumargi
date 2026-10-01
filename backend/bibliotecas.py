# -*- coding: utf-8 -*-
"""Directorio de bibliotecas de la red, leído de datos/bibliotecas.xlsx.

Columnas esperadas (la cabecera puede ir en cualquier orden):

    Nombre_biblioteca   texto, obligatorio
    Sucursal            número entero, obligatorio: el código que aparece en
                        la columna «Suc.» de los listados de AbsysNet
    Poblacion           número entero, obligatorio: habitantes atendidos
    Metros_cuadrados    número, opcional: superficie útil de la biblioteca
    Clave               texto, opcional: la clave con la que la biblioteca
                        entra a su último análisis guardado sin subir los
                        ficheros. La pone y la cambia el administrador de la
                        red; vacía = sin acceso con clave.

Se admite también un bibliotecas.csv con las mismas columnas, por si la
biblioteca no tiene Excel instalado.
"""

import csv
import os
import re
import unicodedata


class ErrorDirectorio(Exception):
    pass


COLUMNAS = ("Nombre_biblioteca", "Sucursal", "Poblacion", "Metros_cuadrados")


def _plano(texto):
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return " ".join(texto.lower().replace("ejemplo:", " ").split())


# «Biblioteca Pública Municipal de» y variantes: se quitan para que baste con
# escribir el nombre de la localidad («Monteagudo»)
_PREFIJO = re.compile(r"^(bibliotecas?|bibl\.?)( (publica|municipal|popular|comarcal))*( (de|del|d'))? +")


def nucleo_nombre(texto):
    return _PREFIJO.sub("", _plano(texto)).strip()


def _normaliza_cabecera(valor):
    return str(valor or "").strip().lower().replace(" ", "_").replace("²", "").replace("á", "a")


def _a_entero(valor):
    if valor is None or str(valor).strip() == "":
        return None
    try:
        return int(float(str(valor).replace(".", "").replace(",", ".")))
    except (TypeError, ValueError):
        return None


def _filas_xlsx(ruta):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ErrorDirectorio(
            "Falta la librería openpyxl para leer el Excel. Instálala con "
            "«pip install openpyxl» o usa un bibliotecas.csv.") from exc
    libro = load_workbook(ruta, read_only=True, data_only=True)
    hoja = libro.active
    filas = hoja.iter_rows(values_only=True)
    try:
        cabecera = [_normaliza_cabecera(c) for c in next(filas)]
    except StopIteration:
        raise ErrorDirectorio("El fichero de bibliotecas está vacío.")
    for fila in filas:
        yield dict(zip(cabecera, fila))
    libro.close()


def _filas_csv(ruta):
    with open(ruta, "r", encoding="utf-8-sig", newline="") as f:
        muestra = f.read(4096)
        f.seek(0)
        try:
            dialecto = csv.Sniffer().sniff(muestra, delimiters=";,\t")
        except csv.Error:
            dialecto = csv.excel
        lector = csv.DictReader(f, dialect=dialecto)
        for fila in lector:
            yield {_normaliza_cabecera(k): v for k, v in fila.items()}


class Directorio:
    """Bibliotecas de la red: nombre, código de sucursal, población y superficie."""

    def __init__(self, ruta):
        self.ruta = ruta
        self.poblacion = {}
        self.superficie = {}
        self.por_sucursal = {}
        self.claves = {}
        self.avisos = []
        self.fecha_fichero = None
        self._cargar()

    def _cargar(self):
        ruta = self.ruta
        if not os.path.exists(ruta):
            alternativa = os.path.splitext(ruta)[0] + ".csv"
            if os.path.exists(alternativa):
                ruta = alternativa
            else:
                raise ErrorDirectorio(
                    f"No se encuentra el fichero de bibliotecas en «{self.ruta}». "
                    f"Consulta el apartado 2 del README.")

        self.fecha_fichero = os.path.getmtime(ruta)
        filas = _filas_csv(ruta) if ruta.lower().endswith(".csv") else _filas_xlsx(ruta)

        for n, fila in enumerate(filas, start=2):
            nombre = str(fila.get("nombre_biblioteca") or "").strip()
            if not nombre:
                continue
            sucursal = _a_entero(fila.get("sucursal"))
            poblacion = _a_entero(fila.get("poblacion"))
            metros = _a_entero(fila.get("metros_cuadrados"))

            if sucursal is None:
                self.avisos.append(f"Fila {n} ({nombre}): sin código de sucursal, se omite.")
                continue
            if sucursal in self.por_sucursal:
                self.avisos.append(
                    f"Fila {n} ({nombre}): el código de sucursal {sucursal} ya está "
                    f"asignado a «{self.por_sucursal[sucursal]}», se omite.")
                continue

            self.por_sucursal[sucursal] = nombre
            clave = fila.get("clave")
            if isinstance(clave, float) and clave.is_integer():
                clave = int(clave)                  # Excel guarda 1234 como 1234.0
            if clave is not None and str(clave).strip():
                self.claves[nombre] = str(clave).strip()
            self.poblacion[nombre] = poblacion or 0
            self.superficie[nombre] = metros

        if not self.por_sucursal:
            raise ErrorDirectorio(
                f"El fichero «{ruta}» no contiene ninguna biblioteca válida. "
                f"Revisa que las columnas se llamen {', '.join(COLUMNAS)}.")

    def nombre_de_sucursal(self, codigo):
        if codigo is None:
            return None
        return self.por_sucursal.get(int(codigo))

    def cambiado(self):
        """¿Se ha modificado el fichero desde que se leyó?"""
        try:
            ruta = self.ruta if os.path.exists(self.ruta) else os.path.splitext(self.ruta)[0] + ".csv"
            return os.path.getmtime(ruta) != self.fecha_fichero
        except OSError:
            return False

    @property
    def hay_claves(self):
        return bool(self.claves)

    def buscar(self, texto):
        """Nombre oficial de la biblioteca a partir de lo que escribe el
        usuario: «Monteagudo», «Biblioteca de Monteagudo»… Sin tildes ni
        mayúsculas. None si no hay ninguna o hay más de una."""
        buscado = nucleo_nombre(texto)
        if not buscado:
            return None
        encontrados = [n for n in self.por_sucursal.values()
                       if nucleo_nombre(n) == buscado or _plano(n) == _plano(texto)]
        return encontrados[0] if len(set(encontrados)) == 1 else None

    def clave_de(self, nombre):
        return self.claves.get(nombre)

    def habitantes(self, nombre):
        return self.poblacion.get(nombre, 0)

    def metros(self, nombre):
        return self.superficie.get(nombre)

    def __len__(self):
        return len(self.por_sucursal)
